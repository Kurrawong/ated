from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("ATED_Options_Paper_Draft.docx")

INK = "243447"
BLUE = "165A7A"
LIGHT_BLUE = "EAF2F6"
LIGHT_GREY = "F2F4F7"
MID_GREY = "667085"
GREEN = "246B49"
AMBER = "8A5A00"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MID_GREY)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after, color in (
        ("Heading 1", 16, 16, 7, BLUE),
        ("Heading 2", 13, 11, 5, BLUE),
        ("Heading 3", 11, 8, 3, INK),
    ):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.46)
        style.paragraph_format.first_line_indent = Inches(-0.22)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08

    header = section.header.paragraphs[0]
    header.text = "ATED SEMANTIC WEB  |  OPTIONS PAPER"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.runs[0], size=8.5, bold=True, color=MID_GREY)
    p_pr = header._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "D0D5DD")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Draft for discussion  •  ")
    set_font(r, size=9, color=MID_GREY)
    add_field(footer, "PAGE")


def para(doc, text="", bold_lead=None, italic=False, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, italic=italic)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.72)
        p.paragraph_format.first_line_indent = Inches(-0.2)
    set_font(p.add_run(text))
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(text))
    return p


def callout(doc, label, text, fill=LIGHT_BLUE, label_color=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label.upper() + "  ")
    set_font(r, size=9, bold=True, color=label_color)
    r = p.add_run(text)
    set_font(r, size=10.5, bold=True, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(doc, headers, rows, widths, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    set_table_geometry(t, widths)
    set_repeat_table_header(t.rows[0])
    for i, value in enumerate(headers):
        cell = t.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_GREY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), size=font_size, bold=True, color=INK)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            set_font(p.add_run(value), size=font_size, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    style_document(doc)
    core = doc.core_properties
    core.title = "Options Paper for Publication and Management of a Semantic Web Version of ATED"
    core.subject = "Next-stage delivery options for ACER"
    core.author = "KurrawongAI"
    core.keywords = "ATED, ACER, Semantic Web, SKOS, Linked Data, Prez"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("OPTIONS PAPER")
    set_font(r, size=11, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Publication and management of a\nSemantic Web version of ATED")
    set_font(r, size=29, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Australian Thesaurus of Education Descriptors")
    set_font(r, size=15, color=BLUE)
    callout(
        doc,
        "Purpose",
        "To select a sustainable next-stage model for publishing, maintaining and governing ATED as reusable Semantic Web data.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    for label, value in (
        ("Prepared for", "Australian Council for Educational Research (ACER)"),
        ("Prepared by", "KurrawongAI"),
        ("Status", "Draft for discussion"),
        ("Date", "3 August 2026"),
    ):
        line = doc.add_paragraph()
        line.paragraph_format.space_after = Pt(4)
        set_font(line.add_run(f"{label}: "), size=10, bold=True, color=MID_GREY)
        set_font(line.add_run(value), size=10, color=INK)

    page_break(doc)

    doc.add_heading("Executive summary", level=1)
    para(doc, "The first project stage has produced a standards-based Semantic Web representation of ATED and its Subject Categories, test-published through KurrawongAI’s demonstration Prez service. The main vocabulary contains 5,280 SKOS concepts and the supporting scheme contains 41 subject categories. ATED’s legacy MultiTes Term Numbers (TNRs) are retained both as concept identifiers and typed notations, preserving a traceable link to the source system.")
    para(doc, "The next stage should make this work operational rather than redesign the vocabulary. The key decisions are where the public service is hosted, how releases are governed, which system remains authoritative for editorial change, and which user-experience improvements are essential for launch.")
    callout(doc, "Recommended direction", "Commission a scoped next-stage implementation of a shared-tenancy vocabulary service: K‑Maker for authenticated editing, validation, review and release management; Prez/Prez UI for read-only public delivery; and a platform layer for tenant isolation and operations. Retain MultiTes as the initial source baseline until the future editorial authority and migration process are accepted.", fill="E8F3ED", label_color=GREEN)
    doc.add_heading("Decisions requested", level=2)
    numbered(doc, "Endorse a managed shared-tenancy Prez service as the preferred initial production model, subject to an agreed service scope, tenancy controls and quotation.")
    numbered(doc, "Confirm the intended management model: MultiTes-led releases initially, or a controlled transition to K‑Maker as the operational editorial authority. Do not permit unmanaged editing in both systems.")
    numbered(doc, "Approve a short production-readiness phase to reconcile the generators with the release files, complete metadata decisions, add automated quality gates and package the formal data handover.")
    numbered(doc, "Prioritise a minimum launch set from ACER’s interface feedback, while managing Prez-dependent enhancements through a transparent backlog.")

    doc.add_heading("Options at a glance", level=2)
    table(doc,
          ["Option", "What ACER receives", "Operational profile", "Assessment"],
          [
              ("1. File handover only", "Validated RDF release and documentation", "No live service", "Lowest commitment; does not realise public discovery or API benefits"),
              ("2. Static web publication", "Files plus generated HTML on ordinary web hosting", "Low-cost, limited search/API capability", "Useful interim or fallback, but weaker user experience"),
              ("3. Shared vocabulary platform", "K‑Maker lifecycle plus public Prez UI/API and tenant operations", "KurrawongAI-operated with ACER governance", "Recommended next-stage design"),
              ("4. Dedicated platform", "Dedicated Prez stack, provider- or ACER-operated", "Separate infrastructure and operations", "Higher isolation/control; higher cost and support burden"),
          ], [1600, 2750, 2450, 2560], font_size=8.5)

    doc.add_heading("1. Purpose and scope", level=1)
    para(doc, "This paper considers how ACER could publish and manage ATED after completion of the current proof-of-concept work. It addresses both the public delivery service and the editorial/governance workflow behind it. These are related but separate decisions: ACER can change hosting without changing the vocabulary identifiers, and can retain MultiTes initially or transition editorial work to K‑Maker while publishing approved RDF through Prez.")
    para(doc, "The paper is based on the current ATED repository, validation of the release files, ACER feedback recorded in GitHub issues 3–11, and the proposed roles of K‑Maker, Prez and Prez UI. It does not imply that ACER production hosting or SaaS work has already been authorised, or that a complete multi-tenant product has already been evidenced publicly. A binding quotation and final security design require discovery and confirmation of ACER’s service levels, identity environment, procurement requirements and expected usage.")

    doc.add_heading("2. Deliverables completed to date", level=1)
    table(doc,
          ["Deliverable", "Current position"],
          [
              ("Semantic Web version", "Completed as SKOS/RDF (Turtle), with stable TNR-based concept IRIs, preferred and alternative labels, definitions/notes, hierarchy, related terms, dates, subjects and source links."),
              ("Test publication", "Completed on the KurrawongAI demonstration Prez environment for ACER review. This is a test environment, not a production service commitment."),
              ("Reusable data-file handover", "Substantively ready. A formal release package and validation statement should now be produced so the handover is explicit, versioned and independently usable."),
              ("Options Paper", "This document provides the decision framework and proposed next-stage design."),
          ], [2300, 7060], font_size=9)

    doc.add_heading("Current data readiness", level=2)
    para(doc, "Automated syntax checks confirm that both Turtle files are valid RDF. The current release contains:")
    bullet(doc, "5,280 ATED concepts and 117,164 RDF statements in the main vocabulary;")
    bullet(doc, "one preferred English label and one typed TNR notation for every ATED concept;")
    bullet(doc, "a source MultiTes URL for every ATED concept;")
    bullet(doc, "subject-category assignments for 5,274 concepts, leaving six concepts to confirm as intentionally uncategorised or correct before release; and")
    bullet(doc, "41 subject-category concepts and 297 RDF statements in the supporting vocabulary.")
    callout(doc, "Release caveat", "The checked-in RDF is valid, but the documented generators do not yet reproduce the release graphs exactly. Recent source-link and scheme-metadata changes were applied after generation. Reconcile the scripts and files before issuing the formal validation statement.", fill="FFF4E5", label_color=AMBER)

    doc.add_heading("Formal handover package for deliverable 3", level=2)
    para(doc, "The handover should be a named, immutable release rather than an informal pointer to a repository. It should contain:")
    bullet(doc, "`ated.ttl` and `ated-sc.ttl`, with a release date and semantic version;")
    bullet(doc, "a plain-language README/data dictionary describing IRIs, SKOS properties, date types, provenance and the relationship between the two schemes;")
    bullet(doc, "a validation report recording syntax, counts, modelling checks, known exceptions and tool versions;")
    bullet(doc, "SHA-256 checksums, a repository tag/commit identifier and the source-extract identifier used to build the release;")
    bullet(doc, "the generation scripts and repeatable commands; and")
    bullet(doc, "an explicit licence and rights statement covering both data reuse and public publication.")

    doc.add_heading("3. Principles for the next stage", level=1)
    table(doc,
          ["Principle", "Design implication"],
          [
              ("ACER retains control", "ACER owns the vocabulary content, release files, identifiers, governance decisions and export history."),
              ("Stable identifiers", "The `https://linked.data.gov.au/def/ated/{TNR}` identifiers remain unchanged when hosting or UI changes."),
              ("Open, standards-based delivery", "Publish SKOS/RDF and an API alongside a usable HTML interface; avoid dependence on a proprietary storage format."),
              ("One editorial authority", "Keep one system of record for content changes. Avoid simultaneous editing in MultiTes and RDF without an explicit reconciliation process."),
              ("Releases are governed", "Separate draft/editorial state from approved public releases, with named owners, review and rollback."),
              ("Portability by design", "Automated builds, manifests, backups and documentation make provider transition practical."),
              ("Public by default, protected administration", "Vocabulary reading and machine access are public; publishing and administration require controlled access."),
          ], [2100, 7260], font_size=9)

    doc.add_heading("4. Publication options", level=1)
    doc.add_heading("Option 1 — validated file handover only", level=2)
    para(doc, "ACER receives a certified RDF release and maintains or redistributes it through existing channels. There is no maintained Semantic Web user interface or live API.")
    bullet(doc, "Advantages: minimal ongoing cost; maximum simplicity; satisfies archival and machine-readable handover.")
    bullet(doc, "Limitations: no resolvable concept service, full-text search, browsing interface, content negotiation or managed release endpoint.")
    bullet(doc, "Best fit: contractual completion or fallback if no hosting is approved.")

    doc.add_heading("Option 2 — static web publication", level=2)
    para(doc, "Publish versioned RDF files and generated HTML through ACER’s ordinary web hosting or a static site. Persistent identifiers redirect to the static pages or files.")
    bullet(doc, "Advantages: low operational complexity; strong portability; easy caching and download.")
    bullet(doc, "Limitations: search and API behaviour are basic; browsing a large thesaurus is less effective; updates still need a release pipeline.")
    bullet(doc, "Best fit: a low-cost public baseline or continuity service.")

    doc.add_heading("Option 3 — KurrawongAI shared vocabulary platform", level=2)
    para(doc, "KurrawongAI implements and operates a small ACER tenancy across K‑Maker and Prez. K‑Maker supports authenticated editing, validation, governance, review and release management. Prez provides the read-only Linked Data API, and Prez UI provides the ACER-branded public experience. The platform layer supplies identity integration, tenant isolation, deployment, jobs, search indexes, monitoring, backups and support.")
    bullet(doc, "Advantages: covers the full maintenance lifecycle, not only hosting; lower unit cost than a dedicated stack; access to vocabulary/Prez expertise; common platform improvements can be incorporated; lower ACER operational burden.")
    bullet(doc, "Limitations: this is a materially broader tenancy problem than shared public hosting. Drafts, permissions, caches, jobs, logs, indexes, configuration, backups and releases all need tested isolation. Product readiness and scope must be established during discovery.")
    bullet(doc, "Best fit: ATED is a bounded, valuable vocabulary with a small editor/reviewer group and public plus machine-readable outputs—strong characteristics for a small shared-platform tenant.")

    doc.add_heading("Option 4 — dedicated Prez platform", level=2)
    para(doc, "Deploy a dedicated RDF database, Prez API and UI for ACER. KurrawongAI could operate the dedicated stack, or it could be deployed into ACER-controlled cloud or on-premises infrastructure and operated by ACER or its infrastructure partner.")
    bullet(doc, "Advantages: strongest workload and infrastructure isolation; greater freedom for ACER-specific configuration; can align directly with ACER security, monitoring and identity standards.")
    bullet(doc, "Limitations: higher implementation and recurring cost; an ACER-operated variant also requires platform engineering, RDF/SPARQL capability, upgrades, incident response, backups and application support.")
    bullet(doc, "Best fit: a later transition if demand, security classification, customisation or service-level requirements justify separate infrastructure.")

    doc.add_heading("5. Comparative assessment", level=1)
    table(doc,
          ["Criterion", "1 File", "2 Static", "3 Shared", "4 Dedicated"],
          [
              ("Public human discovery", "Low", "Medium", "High", "High"),
              ("Machine API / Linked Data", "Low", "Medium", "High", "High"),
              ("Time to production", "Immediate", "Short", "Short–medium", "Medium–long"),
              ("ACER operational effort", "Low", "Low", "Low", "Low–high*"),
              ("Recurring cost", "Very low", "Low", "Medium", "High"),
              ("Control / isolation", "N/A", "Medium–high", "Medium", "High"),
              ("Portability", "High", "High", "High if contracted", "High"),
              ("Ability to address UI feedback", "None", "Low", "High", "High"),
              ("Overall fit for next stage", "Fallback", "Interim", "Preferred", "Future option"),
          ], [2350, 1400, 1500, 2050, 2060], font_size=8.3)
    para(doc, "Ratings are qualitative and should be revisited after ACER confirms availability, security, support and budget requirements. *ACER effort for a dedicated platform is low if KurrawongAI operates it and high if ACER operates it. Hosting cost should be quoted against an agreed service catalogue rather than estimated from the proof of concept.", italic=True)

    doc.add_heading("6. Recommended target design", level=1)
    callout(doc, "Target operating model", "ACER governs content and approves releases. K‑Maker provides the authenticated draft-to-release workflow; Prez and Prez UI publish only approved data. The platform layer isolates and operates the ACER tenancy. MultiTes supplies the initial baseline until ACER accepts which system is authoritative for ongoing edits.")
    doc.add_heading("Logical architecture", level=2)
    para(doc, "The design is deliberately modular. Each layer can be replaced without changing ATED concept identifiers:")
    numbered(doc, "Baseline/import — an approved MultiTes extract is converted by versioned scripts to TNR-based SKOS/RDF and imported as the initial managed release.")
    numbered(doc, "Authoring — named ACER editors work in their K‑Maker tenant; drafts remain separate from the public release.")
    numbered(doc, "Quality gate — K‑Maker and automated pipelines apply syntax, SKOS/profile, identifier, relationship, count and regression checks.")
    numbered(doc, "Review and approval — reviewer/publisher roles inspect the change set and authorise a named release.")
    numbered(doc, "Publication — only approved release data is deployed to the ACER tenant’s Prez named graphs, search index and public presentation.")
    numbered(doc, "Delivery — Prez provides read-only machine access; Prez UI provides ACER-branded search, browse and concept pages.")
    numbered(doc, "Persistence — `linked.data.gov.au` identifiers redirect to the active production presentation while remaining independent of the host.")
    numbered(doc, "Operations — the platform layer controls identity, isolation, jobs, monitoring, backup/restore, upgrades and rollback.")

    doc.add_heading("Access and identity", level=2)
    bullet(doc, "Public read access should not require authentication; this supports indexing, linking and machine reuse.")
    bullet(doc, "Administrative and deployment access should use named accounts, least privilege and multi-factor authentication.")
    bullet(doc, "If ACER requires single sign-on, confirm its corporate OIDC or SAML capability during discovery. Authentication should protect administration, not become a dependency of concept IRIs.")
    bullet(doc, "If a future RDF-native editing interface is introduced, editor, reviewer and publisher roles should be separated.")

    doc.add_heading("Shared-tenancy controls", level=2)
    para(doc, "The service definition should make shared infrastructure transparent without implying shared governance or uncontrolled data access. At minimum it should specify:")
    bullet(doc, "logical separation of ACER catalogue, content and supporting graphs, with tenant-scoped deployment permissions;")
    bullet(doc, "ACER-specific catalogue identity, public route/domain, branding, profiles and search configuration where supported;")
    bullet(doc, "capacity controls and monitoring that protect ATED from another tenant’s workload;")
    bullet(doc, "backup, restore and rollback procedures capable of restoring ACER data to a selected approved release;")
    bullet(doc, "tenant-aware incident reporting, usage/service reporting and deletion/exit procedures; and")
    bullet(doc, "clear classification of common platform enhancements versus ACER-funded configuration or feature development.")

    doc.add_heading("Management model", level=2)
    para(doc, "ACER’s likely ongoing requirement includes editing, validation, review, publication and maintenance. The next stage should therefore test K‑Maker as the tenant-aware management layer, with Prez remaining a read-only publisher. The project must still make an explicit source-of-truth decision: either MultiTes continues to lead and K‑Maker receives controlled imports, or editorial authority transitions to K‑Maker after migration and acceptance. Independent editing in both systems is not sustainable.")
    para(doc, "The K‑Maker training/demo environment being prepared with ATED as an example is a useful precursor: it can exercise temporary access, draft/public separation, validation, review/publication, reset/restore and operator documentation. It is not itself an ACER production system and should not be represented as proof that production tenancy is complete.")

    doc.add_heading("Requirement classification", level=2)
    para(doc, "Each next-stage requirement should be classified before estimation to control cost and avoid client-specific forks:")
    numbered(doc, "Core Prez/K‑Maker capability — behaviour that belongs in the underlying products.")
    numbered(doc, "Reusable platform feature — tenancy, identity, isolation, deployment or operations useful across customers.")
    numbered(doc, "ACER configuration/theme — labels, field ordering, visual identity, profiles and tenant policy.")
    numbered(doc, "ACER-only extension — a consciously costed feature used only when configuration or reusable product work cannot meet the requirement.")

    doc.add_heading("7. ACER feedback and proposed treatment", level=1)
    para(doc, "The open repository issues reflect valuable product requirements. They should be handled as a prioritised service backlog, with the distinction below made clear: data/model tasks are controlled within the ATED project; presentation tasks require Prez UI configuration or development; and search tasks may require Prez/API query changes.")
    table(doc,
          ["Issue", "Requirement", "Proposed treatment", "Dependency / priority"],
          [
              ("#3", "Subject Category provenance and metadata", "Confirm ownership, dates, definition, independent use and mappings; then complete scheme metadata.", "ATED/ACER decision • pre-release"),
              ("#4", "Duplicate search results", "Group matches by concept and use weighted scoring; validate the existing weighted-search candidate.", "Prez search change • launch priority"),
              ("#5", "Alphabetical search order", "Use relevance for search by default, with an alphabetical sort/browse option. Pure alphabetical search can bury exact matches.", "Product decision + Prez • high"),
              ("#6", "Demote scheme name", "Retain as secondary context, visually quieter; essential when multiple schemes are searched.", "Prez UI • medium"),
              ("#7", "Show ‘matched on’", "Display the matched label/property directly in each result, especially when an alternative label caused the hit.", "Prez UI/search payload • high"),
              ("#8", "Reorder concept fields", "Use an ATED concept profile: scope/definition, preferred label, broader/narrower/related, alternatives, subject, dates, notation/source.", "Prez profile/UI • launch priority"),
              ("#9", "Source IRIs", "Implemented in the current RDF as `dcterms:source` using TNR-based MultiTes URLs; verify display and close after release validation.", "Data complete; UI check"),
              ("#10", "Alphabetise relation values", "Sort displayed labels alphabetically within broader, narrower and related groups.", "Prez API/UI • medium"),
              ("#11", "Top-align field names", "Apply top alignment for multi-value property rows.", "Prez UI • medium"),
          ], [800, 1950, 4200, 2410], font_size=7.8)

    doc.add_heading("Minimum launch experience", level=2)
    para(doc, "A production launch need not wait for every general Prez enhancement. The minimum acceptable ATED experience should include:")
    bullet(doc, "one result per concept, with exact preferred-label matches ranked strongly;")
    bullet(doc, "clear indication when a result matched an alternative label;")
    bullet(doc, "a concept page ordered around scope, preferred form and semantic relationships;")
    bullet(doc, "stable links to all concepts and downloadable RDF;")
    bullet(doc, "alphabetical browse/sort and readable multi-value fields; and")
    bullet(doc, "an accessible, ACER-approved visual treatment and service information page.")

    doc.add_heading("8. Delivery plan", level=1)
    table(doc,
          ["Phase", "Indicative scope", "Exit outcome"],
          [
              ("0. Decision and discovery", "Confirm preferred option, K‑Maker/Prez product boundary, source of truth, tenancy readiness, stakeholders, service levels, identity, domains/PIDs, licence, accessibility and procurement constraints.", "Approved architecture, scope and commercial proposal"),
              ("1. Production-ready data release", "Reconcile generators; review six uncategorised concepts; complete scheme metadata; add automated validation and regression checks; tag and package the handover.", "ATED Semantic Web Release 1.0 with validation statement"),
              ("2. Shared-platform implementation", "Provision ACER tenancy; configure K‑Maker roles/workflows and draft/public separation; deploy Prez publication, search and branding; tenant isolation, security and accessibility testing.", "Production candidate ready for acceptance"),
              ("3. Acceptance and launch", "ACER user acceptance, defect remediation, operational runbook, training, support and rollback test.", "Public launch and accepted service"),
              ("4. Operate and improve", "Scheduled releases, monitoring, backups, service reporting and prioritised enhancement backlog.", "Sustainable operating cadence"),
              ("5. Strategic review", "Assess usage, editorial needs, platform performance and whether to retain managed hosting, transition to ACER, or adopt RDF-native management.", "Evidence-based renewal or transition decision"),
          ], [1750, 5150, 2460], font_size=8.4)

    doc.add_heading("Indicative responsibilities", level=2)
    table(doc,
          ["Activity", "ACER", "Service provider"],
          [
              ("Vocabulary content and editorial policy", "Accountable; edit, review and approve", "Provide K‑Maker workflow; advise on modelling and quality"),
              ("Source export and release request", "Provide/authorise", "Receive and verify"),
              ("K‑Maker, transformation and validation", "Use workflow; review evidence", "Implement, maintain and operate tenant capability"),
              ("Publication approval", "Accountable", "Execute approved release"),
              ("Shared platform hosting, monitoring, backup and recovery", "Set requirements; receive tenant-level reports", "Operate under agreed service levels and tenancy controls"),
              ("Platform enhancement priorities", "Prioritise and accept", "Estimate, develop/configure and test"),
              ("Persistent identifier governance", "Own identifier policy", "Implement/maintain redirect target as agreed"),
          ], [3200, 3100, 3060], font_size=8.6)

    doc.add_heading("Commercial scoping inputs", level=2)
    para(doc, "A quotation should separate one-off implementation from recurring service and optional enhancement work. ACER should provide or agree:")
    bullet(doc, "availability and support hours, recovery targets and planned maintenance expectations;")
    bullet(doc, "anticipated public traffic, API use and any rate-limiting requirements;")
    bullet(doc, "ACER domain/branding, accessibility and analytics requirements;")
    bullet(doc, "release frequency and expected scale of vocabulary changes;")
    bullet(doc, "security assurance, data residency, logging and identity requirements; and")
    bullet(doc, "the boundary between included platform updates and separately funded ATED-specific enhancements.")

    doc.add_heading("9. Risks and controls", level=1)
    table(doc,
          ["Risk", "Consequence", "Primary control"],
          [
              ("Generator/release drift", "A future rebuild changes or loses approved data", "Make generators authoritative; graph-level regression tests; immutable releases"),
              ("Dual editing", "MultiTes and RDF diverge", "One editorial source of truth; no direct production RDF edits"),
              ("Persistent IRI failure", "External links break when hosting changes", "Maintain PID allocation/redirects separately from the application host"),
              ("Prez enhancement dependency", "UX requirements delay launch", "Define minimum launch set; configuration first; maintain staged backlog and fallbacks"),
              ("Provider or shared-platform lock-in", "Transition becomes expensive", "ACER-owned RDF, Git history, manifests, tenant exports, runbook and exit obligations"),
              ("Shared-tenancy isolation failure", "Another tenant affects drafts, permissions, availability, confidentiality or recovery", "Isolation tests covering data, caches, jobs, logs, indexes, configuration, backups and releases"),
              ("Unclear metadata rights", "Publication or reuse uncertainty", "Approve licence, attribution, creator/publisher and source statements before release"),
              ("Unsupported operations", "Outages, stale data or security exposure", "Named service owner, monitoring, patching, backups, restore tests and incident process"),
              ("Search behaviour disappoints users", "Low adoption and continued use of legacy interface", "ACER acceptance scenarios, analytics and iterative tuning"),
          ], [2350, 2940, 4070], font_size=8.4)

    doc.add_heading("10. Recommendation", level=1)
    para(doc, "Proceed to scope Option 3, KurrawongAI’s proposed shared vocabulary platform, for an initial fixed production period. This offers the clearest route from a successful proof of concept to a maintainable service: K‑Maker supports ACER’s editorial lifecycle, while Prez/Prez UI deliver approved data publicly. It also spreads common platform operating costs across tenants. The next-stage engagement should validate tenancy and product readiness before any production commitment.")
    para(doc, "Before production publication, complete a short release-readiness package: reconcile generation scripts and checked-in RDF; confirm the six missing subject assignments; settle Subject Category metadata and licensing; automate validation; and issue a versioned handover with checksums and a validation statement. These are controlled, finite tasks and should not be conflated with broader Prez product enhancements.")
    para(doc, "Adopt the proposed minimum launch experience and test it with ACER users. Treat weighted/deduplicated search, visible match context and concept-field ordering as launch priorities. Retain other presentation refinements in the next-stage backlog where they depend on Prez development.")

    doc.add_heading("Proposed next decision", level=2)
    callout(doc, "Decision", "Authorise a scoped discovery and production-readiness phase that ends with: (1) ATED Semantic Web Release 1.0 and formal handover; (2) an accepted K‑Maker/Prez shared-tenancy architecture and source-of-truth decision; (3) a service definition and costed implementation proposal; and (4) an agreed launch backlog mapped to ACER feedback.", fill="E8F3ED", label_color=GREEN)

    doc.add_heading("Assumptions requiring ACER confirmation", level=2)
    bullet(doc, "ATED and the Subject Categories may be published publicly under an agreed reuse licence.")
    bullet(doc, "MultiTes remains available and authoritative for editorial change during the initial production period.")
    bullet(doc, "Public read/API access is desirable and does not require user authentication.")
    bullet(doc, "ACER can nominate content owner, technical contact and publication approver roles.")
    bullet(doc, "The `linked.data.gov.au` ATED namespace and redirects can be governed for long-term persistence.")

    doc.add_heading("Appendix A — Release validation statement template", level=1)
    table(doc,
          ["Field", "Release record"],
          [
              ("Release name/version", "ATED Semantic Web Release [version]"),
              ("Release date", "[date]"),
              ("Source extract", "[MultiTes export filename/date/checksum]"),
              ("Repository state", "[Git commit and signed/annotated tag]"),
              ("Files", "ated.ttl; ated-sc.ttl; documentation; validation report"),
              ("RDF syntax", "RIOT and Raptor parse without error"),
              ("Profile/model checks", "[SKOS/SHACL and project invariant results]"),
              ("Counts", "[concepts, triples, labels, relationships, categories]"),
              ("Known exceptions", "[documented exceptions, including uncategorised concepts if accepted]"),
              ("Checksums", "[SHA-256 for each delivered file]"),
              ("Licence and attribution", "[approved statement]"),
              ("Approval", "[ACER approver and date]"),
          ], [2400, 6960], font_size=9)
    para(doc, "Suggested certification wording: “The listed files were generated from the identified ATED source extract using the supplied versioned transformation code. They parse as valid RDF/Turtle and passed the checks recorded in the accompanying validation report. Known exceptions are stated explicitly. The files are self-contained publication artefacts and do not depend on continued hosting by KurrawongAI.”")

    doc.add_heading("Appendix B — Evidence and references", level=1)
    sources = [
        "ATED repository and issue register, Kurrawong/ated, reviewed 3 August 2026: https://github.com/Kurrawong/ated",
        "ATED ACER feedback issues 3–11: https://github.com/Kurrawong/ated/issues",
        "W3C, SKOS Simple Knowledge Organization System Reference: https://www.w3.org/TR/skos-reference/",
        "KurrawongAI, Prez Overview and deployment/data-management documentation: https://docs.kurrawong.ai/products/prez/",
        "KurrawongAI, Prez Manifest Model: https://docs.kurrawong.ai/products/prez/manifest/",
        "Kurrawong internal training/demo issue #96, ‘Set up K‑Maker for training demonstrations’ (contextual precursor; not an ACER production authorisation).",
        "Australian Government Linked Data Working Group, persistent identifier services and governance: https://www.linked.data.gov.au/",
    ]
    for s in sources:
        bullet(doc, s)
    para(doc, "Repository validation performed 3 August 2026 using RIOT, Raptor and Kurra. Counts in this paper describe commit 4696757 on the local `main` branch; the uncommitted `AGENTS.md` project-guidance file does not affect the vocabulary data.", italic=True)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
