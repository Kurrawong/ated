from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path("/Users/leskneebone/Downloads/ATED_Options_Paper_Edited.docx")
OUTPUT = Path(__file__).with_name("ATED_Options_Paper_with_Appendix_C.docx")

CONCEPTS = [
    ("Brain drain", "10817", "2017-02", "Employment problems; Migration patterns"),
    ("Foetal alcohol syndrome", "7262", "Not supplied", "Not supplied"),
    ("Man machine interface", "8606", "Not supplied", "Not supplied"),
    ("Randomised controlled trials", "10903", "2018-08", "Trials"),
    ("Teaching persistence", "9969", "Not supplied", "Not supplied"),
    ("Torres Strait Islander health", "10852", "Not supplied", "Not supplied"),
]


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])


def copy_cell_format(source, destination):
    source_pr = source._tc.tcPr
    destination_pr = destination._tc.get_or_add_tcPr()
    for child in list(destination_pr):
        destination_pr.remove(child)
    for child in source_pr:
        destination_pr.append(deepcopy(child))


def fill_cell(cell, text, paragraph_source):
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = paragraph_source.style
    p.paragraph_format.space_before = paragraph_source.paragraph_format.space_before
    p.paragraph_format.space_after = paragraph_source.paragraph_format.space_after
    p.paragraph_format.line_spacing = paragraph_source.paragraph_format.line_spacing
    r = p.add_run(text)
    if paragraph_source.runs:
        source_rpr = paragraph_source.runs[0]._r.rPr
        if source_rpr is not None:
            r._r.insert(0, deepcopy(source_rpr))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def main():
    doc = Document(SOURCE)
    template = doc.tables[-1]
    header_source = template.rows[0].cells[0]
    body_source = template.rows[1].cells[0]

    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    heading = doc.add_paragraph(style="Heading")
    heading.add_run("Appendix C — ATED concepts without subject categories")

    intro = doc.add_paragraph(style="Body")
    intro.add_run(
        "The following six concepts have no dcterms:subject relationship in the "
        "current ATED RDF release. Each concept is otherwise represented as a valid "
        "SKOS concept with a preferred label, TNR-based IRI and notation, and a "
        "MultiTes source link. Their uncategorised status should be reviewed with "
        "ACER before Release 1.0; it may reflect missing source data or an intentional "
        "absence of classification."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Normal"
    headers = ("Preferred label", "TNR", "Modified", "Broader term(s)")
    widths = [3300, 1100, 1300, 3660]
    for i, text in enumerate(headers):
        copy_cell_format(header_source, table.rows[0].cells[i])
        fill_cell(table.rows[0].cells[i], text, header_source.paragraphs[0])
    set_repeat_header(table.rows[0])

    for label, tnr, modified, broader in CONCEPTS:
        cells = table.add_row().cells
        values = (label, tnr, modified, broader)
        for i, text in enumerate(values):
            copy_cell_format(body_source, cells[i])
            fill_cell(cells[i], text, body_source.paragraphs[0])
    set_table_geometry(table, widths)

    note = doc.add_paragraph(style="Body")
    note.paragraph_format.space_before = doc.paragraphs[-3].paragraph_format.space_before
    run = note.add_run(
        "Concept IRIs are formed as https://linked.data.gov.au/def/ated/{TNR}. "
        "For example, Brain drain is identified by "
        "https://linked.data.gov.au/def/ated/10817. “Not supplied” means that the "
        "corresponding dcterms:modified or skos:broader value is absent from the "
        "current RDF; it is not a validation error."
    )
    if doc.paragraphs[164].runs:
        source_rpr = doc.paragraphs[164].runs[0]._r.rPr
        if source_rpr is not None:
            run._r.insert(0, deepcopy(source_rpr))

    source = doc.add_paragraph(style="Body")
    source_run = source.add_run(
        "Source: vocabs/ated.ttl at repository commit 4696757; checked with Kurra "
        "on 3 August 2026."
    )
    if doc.paragraphs[164].runs:
        source_rpr = doc.paragraphs[164].runs[0]._r.rPr
        if source_rpr is not None:
            source_run._r.insert(0, deepcopy(source_rpr))
    source_run.italic = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
